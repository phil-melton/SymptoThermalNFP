"""Tests for the Flask web UI: pages, forms, Excel upload/download, JSON APIs."""

from __future__ import annotations

import datetime as dt
import io

import pytest
from openpyxl import load_workbook

from cycle_builders import textbook_cycle
from symptothermal_nfp.excel_store import ExcelStore, build_workbook
from symptothermal_nfp.models import AppSettings
from symptothermal_nfp.web import create_app

START = dt.date(2026, 3, 1)


@pytest.fixture()
def data_path(tmp_path):
    return tmp_path / "data.xlsx"


@pytest.fixture()
def client(data_path):
    app = create_app(data_path)
    app.config["TESTING"] = True
    return app.test_client()


def _upload(client, payload: bytes, *, mode: str = "merge", filename: str = "upload.xlsx"):
    return client.post(
        "/upload",
        data={"workbook": (io.BytesIO(payload), filename), "mode": mode},
        content_type="multipart/form-data",
        follow_redirects=True,
    )


def _textbook_workbook_bytes() -> bytes:
    buffer = io.BytesIO()
    build_workbook(textbook_cycle(START), AppSettings()).save(buffer)
    return buffer.getvalue()


def test_index_renders_empty_state(client) -> None:
    response = client.get("/")
    assert response.status_code == 200
    page = response.get_data(as_text=True)
    assert "SymptoThermalNFP" in page
    assert "No observations yet" in page


def test_log_observation_via_form(client, data_path) -> None:
    response = client.post(
        "/observations",
        data={
            "date": "2026-03-01",
            "temperature": "36.55",
            "temperature_unit": "celsius",
            "temperature_time": "06:30",
            "fluid_sensation": "watery",
            "fluid_quantity": "high",
            "fluid_color": "clear",
            "fluid_texture": "stretchy",
            "bleeding": "none",
            "notes": "web test",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    page = response.get_data(as_text=True)
    assert "Logged observation for 2026-03-01" in page
    assert "web test" in page

    stored = ExcelStore(data_path).load_observations()
    assert len(stored) == 1
    assert stored[0].waking_temperature == 36.55
    assert stored[0].fluid is not None
    assert stored[0].fluid.sensation.value == "watery"


def test_log_observation_form_updates_existing_day(client, data_path) -> None:
    for temp in ("36.40", "36.80"):
        client.post(
            "/observations",
            data={"date": "2026-03-01", "temperature": temp, "temperature_unit": "celsius", "bleeding": "none"},
            follow_redirects=True,
        )
    stored = ExcelStore(data_path).load_observations()
    assert len(stored) == 1
    assert stored[0].waking_temperature == 36.8


def test_invalid_form_shows_error_and_saves_nothing(client, data_path) -> None:
    response = client.post(
        "/observations",
        data={"date": "2026-03-01", "temperature": "abc", "bleeding": "none"},
        follow_redirects=True,
    )
    assert "Could not save observation" in response.get_data(as_text=True)
    assert ExcelStore(data_path).load_observations() == []


def test_delete_observation(client, data_path) -> None:
    client.post(
        "/observations",
        data={"date": "2026-03-01", "temperature": "36.5", "temperature_unit": "celsius", "bleeding": "none"},
    )
    response = client.post("/observations/2026-03-01/delete", follow_redirects=True)
    assert "Deleted observation for 2026-03-01" in response.get_data(as_text=True)
    assert ExcelStore(data_path).load_observations() == []

    response = client.post("/observations/2026-03-01/delete", follow_redirects=True)
    assert "No observation found" in response.get_data(as_text=True)


def test_settings_form_round_trip(client, data_path) -> None:
    response = client.post(
        "/settings",
        data={
            "temperature_unit": "fahrenheit",
            "default_wake_time": "05:45",
            "rule_context": "postpartum",
            "transition_cycle_count": "2",
            "track_cervical_position": "on",
            "bip_enabled": "on",
        },
        follow_redirects=True,
    )
    assert "Settings saved" in response.get_data(as_text=True)
    settings = ExcelStore(data_path).load_settings()
    assert settings.temperature_unit.value == "fahrenheit"
    assert settings.rule_context.value == "postpartum"
    assert settings.transition_cycle_count == 2
    assert settings.track_cervical_position
    assert settings.bip_enabled
    assert not settings.use_doering_rule


def test_upload_excel_merge(client, data_path) -> None:
    response = _upload(client, _textbook_workbook_bytes())
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "Imported 28 observation(s)" in page
    assert len(ExcelStore(data_path).load_observations()) == 28


def test_upload_excel_replace_mode(client, data_path) -> None:
    client.post(
        "/observations",
        data={"date": "2020-01-01", "temperature": "36.5", "temperature_unit": "celsius", "bleeding": "none"},
    )
    _upload(client, _textbook_workbook_bytes(), mode="replace")
    stored = ExcelStore(data_path).load_observations()
    assert len(stored) == 28
    assert stored[0].observation_date == START  # 2020 row is gone


def test_upload_rejects_non_workbook(client, data_path) -> None:
    response = _upload(client, b"definitely not xlsx", filename="fake.xlsx")
    assert "Import failed" in response.get_data(as_text=True)
    assert ExcelStore(data_path).load_observations() == []


def test_upload_without_file_flashes_error(client) -> None:
    response = client.post("/upload", data={"mode": "merge"}, follow_redirects=True)
    assert "Choose an .xlsx file" in response.get_data(as_text=True)


def test_upload_reports_skipped_rows(client, data_path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Observations"
    sheet.append(["date", "temperature", "temperature_unit"])
    sheet.append(["2026-03-01", 36.5, "celsius"])
    sheet.append(["2026-03-02", "oops", "celsius"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    response = _upload(client, buffer.getvalue())
    page = response.get_data(as_text=True)
    assert "Imported 1 observation(s)" in page
    assert "Skipped row 3" in page
    assert len(ExcelStore(data_path).load_observations()) == 1


def test_download_returns_current_workbook(client) -> None:
    client.post(
        "/observations",
        data={"date": "2026-03-01", "temperature": "36.5", "temperature_unit": "celsius", "bleeding": "none"},
    )
    response = client.get("/download")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(io.BytesIO(response.data))
    sheet = workbook["Observations"]
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=1).value == "2026-03-01"


def test_template_download_has_headers_only(client) -> None:
    response = client.get("/template.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    sheet = workbook["Observations"]
    assert sheet.cell(row=1, column=1).value == "date"
    assert sheet.max_row == 1


def test_api_observations(client) -> None:
    _upload(client, _textbook_workbook_bytes())
    response = client.get("/api/observations")
    assert response.status_code == 200
    assert len(response.json) == 28
    assert response.json[0]["observation_date"] == "2026-03-01"


def test_uploaded_excel_drives_fertile_day_interpretation_end_to_end(client) -> None:
    """Full stack: xlsx upload -> Excel store -> rule pack -> JSON API."""
    _upload(client, _textbook_workbook_bytes())

    response = client.get("/api/interpretation")
    assert response.status_code == 200
    report = response.json
    assert report["rule_pack_version"] == "stm-v1"
    assert len(report["cycles"]) == 1

    cycle = report["cycles"][0]
    assert cycle["peak_day"] == "2026-03-14"
    assert cycle["temperature_shift"]["confirmed_date"] == "2026-03-17"
    assert cycle["absolute_infertility_start_date"] == "2026-03-17"

    statuses = {day["cycle_day"]: day["status"] for day in cycle["days"]}
    assert statuses[16] == "potentially_fertile"
    assert statuses[17] == "absolute_infertility_from_evening"
    assert statuses[18] == "absolute_infertility"
    assert statuses[28] == "absolute_infertility"

    # The dashboard reflects the same result.
    page = client.get("/").get_data(as_text=True)
    assert "Infertile (Phase 3)" in page
    assert "2026-03-17" in page


def test_index_renders_chart_payload_and_history(client) -> None:
    _upload(client, _textbook_workbook_bytes())
    page = client.get("/").get_data(as_text=True)
    assert '"points":' in page or "&#34;points&#34;:" in page
    assert "Cycle chart" in page
    assert "Observation history" in page
    assert "2026-03-14" in page


def test_data_file_persists_across_app_restarts(data_path) -> None:
    first = create_app(data_path).test_client()
    first.post(
        "/observations",
        data={"date": "2026-03-01", "temperature": "36.5", "temperature_unit": "celsius", "bleeding": "none"},
    )

    second = create_app(data_path).test_client()
    response = second.get("/api/observations")
    assert len(response.json) == 1
