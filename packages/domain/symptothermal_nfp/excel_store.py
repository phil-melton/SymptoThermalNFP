"""Excel-backed persistence for observations and settings.

The workbook layout is deliberately human-friendly so users can edit or build
files in Excel/LibreOffice and upload them:

* ``Observations`` sheet: one row per day, one column per field.
* ``Settings`` sheet: key/value pairs.
"""

from __future__ import annotations

import io
import os
import shutil
import threading
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, BinaryIO, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import (
    AppSettings,
    CervicalPositionObservation,
    DailyObservation,
    FluidObservation,
    parse_hhmm_time,
    parse_iso_date,
)
from .taxonomy import (
    BleedingLevel,
    CervixFirmness,
    CervixHeight,
    CervixOpening,
    FluidQuantity,
    FluidSensation,
    MucusColor,
    MucusTexture,
    TemperatureUnit,
)

OBSERVATIONS_SHEET = "Observations"
SETTINGS_SHEET = "Settings"

OBSERVATION_COLUMNS: tuple[str, ...] = (
    "date",
    "temperature",
    "temperature_unit",
    "temperature_time",
    "temperature_disturbed",
    "fluid_sensation",
    "fluid_quantity",
    "fluid_color",
    "fluid_texture",
    "fluid_amount",
    "peak_quality",
    "cervix_height",
    "cervix_firmness",
    "cervix_opening",
    "bleeding",
    "notes",
)

# Friendly header aliases so hand-made spreadsheets still import cleanly.
_HEADER_ALIASES: dict[str, str] = {
    "observation_date": "date",
    "day": "date",
    "temp": "temperature",
    "waking_temperature": "temperature",
    "bbt": "temperature",
    "unit": "temperature_unit",
    "temp_unit": "temperature_unit",
    "time": "temperature_time",
    "temp_time": "temperature_time",
    "disturbed": "temperature_disturbed",
    "temp_disturbed": "temperature_disturbed",
    "sensation": "fluid_sensation",
    "mucus_sensation": "fluid_sensation",
    "quantity": "fluid_quantity",
    "mucus_quantity": "fluid_quantity",
    "color": "fluid_color",
    "mucus_color": "fluid_color",
    "texture": "fluid_texture",
    "mucus_texture": "fluid_texture",
    "amount": "fluid_amount",
    "mucus_amount": "fluid_amount",
    "peak": "peak_quality",
    "height": "cervix_height",
    "firmness": "cervix_firmness",
    "opening": "cervix_opening",
    "note": "notes",
}

_TRUE_STRINGS = {"true", "yes", "y", "1", "x"}
_FALSE_STRINGS = {"false", "no", "n", "0", ""}


class ExcelStoreError(Exception):
    """Raised when a workbook cannot be read as observation data."""


@dataclass(slots=True)
class RowError:
    row: int
    message: str

    def as_dict(self) -> dict[str, Any]:
        return {"row": self.row, "message": self.message}


@dataclass(slots=True)
class ImportResult:
    imported: int = 0
    replaced: int = 0
    errors: list[RowError] = field(default_factory=list)
    settings_imported: bool = False

    @property
    def ok(self) -> bool:
        return not self.errors

    def as_dict(self) -> dict[str, Any]:
        return {
            "imported": self.imported,
            "replaced": self.replaced,
            "errors": [item.as_dict() for item in self.errors],
            "settings_imported": self.settings_imported,
        }


class ExcelStore:
    """Excel-file-backed store for observations and settings.

    The .xlsx file at ``path`` is the single source of truth: every write
    rewrites the workbook, so the file stays openable in Excel at all times.

    Writes are crash-safe: each save goes to a temp file that atomically
    replaces the data file, and the previous version is kept as a rolling
    ``.bak`` sibling. Mutations are serialized with an in-process lock so
    concurrent requests cannot interleave read-modify-write cycles.
    """

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        self._lock = threading.RLock()

    @property
    def backup_path(self) -> Path:
        return self.path.with_name(self.path.name + ".bak")

    def initialize(self) -> None:
        with self._lock:
            if not self.path.exists():
                self.path.parent.mkdir(parents=True, exist_ok=True)
                self._write(observations=[], settings=AppSettings())

    def load_observations(self) -> list[DailyObservation]:
        if not self.path.exists():
            return []
        workbook = load_workbook(self.path, data_only=True)
        observations, errors = read_observations(workbook)
        if errors:
            details = "; ".join(f"row {item.row}: {item.message}" for item in errors)
            raise ExcelStoreError(f"Stored workbook contains invalid rows: {details}")
        return observations

    def load_settings(self) -> AppSettings:
        if not self.path.exists():
            return AppSettings()
        workbook = load_workbook(self.path, data_only=True)
        return read_settings(workbook) or AppSettings()

    def save_settings(self, settings: AppSettings) -> None:
        with self._lock:
            self._write(observations=self._existing_observations(), settings=settings)

    def upsert_observation(self, observation: DailyObservation) -> bool:
        """Insert or replace the observation for its date. Returns True if replaced."""
        with self._lock:
            observations = self._existing_observations()
            replaced = any(
                item.observation_date == observation.observation_date for item in observations
            )
            observations = [
                item
                for item in observations
                if item.observation_date != observation.observation_date
            ]
            observations.append(observation)
            self._write(observations=observations, settings=self.load_settings())
            return replaced

    def delete_observation(self, observation_date: date) -> bool:
        with self._lock:
            observations = self._existing_observations()
            remaining = [item for item in observations if item.observation_date != observation_date]
            if len(remaining) == len(observations):
                return False
            self._write(observations=remaining, settings=self.load_settings())
            return True

    def import_workbook(
        self,
        source: str | Path | bytes | BinaryIO,
        *,
        mode: str = "merge",
    ) -> ImportResult:
        """Import an uploaded workbook.

        ``mode='merge'`` upserts uploaded rows into the existing data;
        ``mode='replace'`` discards existing observations first. Rows that
        fail validation are reported and skipped; valid rows still import.
        """
        if mode not in {"merge", "replace"}:
            raise ValueError("Import mode must be 'merge' or 'replace'.")

        workbook = _open_workbook(source)
        incoming, errors = read_observations(workbook)
        incoming_settings = read_settings(workbook)

        with self._lock:
            existing = [] if mode == "replace" else self._existing_observations()
            by_date = {item.observation_date: item for item in existing}
            replaced = sum(1 for item in incoming if item.observation_date in by_date)
            for item in incoming:
                by_date[item.observation_date] = item

            settings = incoming_settings or self.load_settings()
            self._write(observations=list(by_date.values()), settings=settings)
        return ImportResult(
            imported=len(incoming),
            replaced=replaced,
            errors=errors,
            settings_imported=incoming_settings is not None,
        )

    def export_bytes(self) -> bytes:
        self.initialize()
        return self.path.read_bytes()

    def _existing_observations(self) -> list[DailyObservation]:
        return self.load_observations()

    def _write(self, *, observations: Iterable[DailyObservation], settings: AppSettings) -> None:
        workbook = build_workbook(observations, settings)
        temp_path = self.path.with_name(self.path.name + ".tmp")
        workbook.save(temp_path)
        try:
            if self.path.exists():
                shutil.copy2(self.path, self.backup_path)
            os.replace(temp_path, self.path)
        finally:
            temp_path.unlink(missing_ok=True)


def build_workbook(
    observations: Iterable[DailyObservation],
    settings: AppSettings | None = None,
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = OBSERVATIONS_SHEET
    _write_observation_sheet(sheet, observations)
    _write_settings_sheet(workbook.create_sheet(SETTINGS_SHEET), settings or AppSettings())
    return workbook


def read_observations(workbook: Workbook) -> tuple[list[DailyObservation], list[RowError]]:
    sheet = _find_observation_sheet(workbook)
    header_map = _read_header_map(sheet)
    if "date" not in header_map.values():
        raise ExcelStoreError(
            f"The '{sheet.title}' sheet needs a 'date' column in its header row."
        )

    observations: list[DailyObservation] = []
    errors: list[RowError] = []
    seen_dates: set[date] = set()

    for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = {
            header_map[col]: row[col]
            for col in header_map
            if col < len(row) and not _is_blank(row[col])
        }
        if not values:
            continue
        try:
            observation = _row_to_observation(values)
        except (ValueError, TypeError) as exc:
            errors.append(RowError(row=row_index, message=str(exc)))
            continue
        if observation.observation_date in seen_dates:
            errors.append(
                RowError(
                    row=row_index,
                    message=f"Duplicate date {observation.observation_date.isoformat()}.",
                )
            )
            continue
        seen_dates.add(observation.observation_date)
        observations.append(observation)

    observations.sort(key=lambda item: item.observation_date)
    return observations, errors


def read_settings(workbook: Workbook) -> AppSettings | None:
    if SETTINGS_SHEET not in workbook.sheetnames:
        return None
    sheet = workbook[SETTINGS_SHEET]
    raw: dict[str, Any] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or _is_blank(row[0]):
            continue
        key = str(row[0]).strip()
        raw[key] = row[1] if len(row) > 1 else None

    if not raw:
        return None

    defaults = AppSettings()
    try:
        return AppSettings.from_dict(
            {
                "temperature_unit": _string_or(raw.get("temperature_unit"), defaults.temperature_unit.value),
                "default_wake_time": _time_string_or(raw.get("default_wake_time"), defaults.default_wake_time),
                "track_cervical_position": _parse_bool(raw.get("track_cervical_position"), default=False),
                "rule_context": _string_or(raw.get("rule_context"), defaults.rule_context.value),
                "transition_cycle_count": int(raw.get("transition_cycle_count") or 0),
                "use_doering_rule": _parse_bool(raw.get("use_doering_rule"), default=False),
                "use_rotzer_rule": _parse_bool(raw.get("use_rotzer_rule"), default=False),
                "bip_enabled": _parse_bool(raw.get("bip_enabled"), default=False),
            }
        )
    except (ValueError, TypeError) as exc:
        raise ExcelStoreError(f"The '{SETTINGS_SHEET}' sheet is invalid: {exc}") from exc


def _open_workbook(source: str | Path | bytes | BinaryIO) -> Workbook:
    try:
        if isinstance(source, bytes):
            return load_workbook(io.BytesIO(source), data_only=True)
        return load_workbook(source, data_only=True)
    except ExcelStoreError:
        raise
    except Exception as exc:  # openpyxl raises several unrelated types here
        raise ExcelStoreError(
            "The file could not be opened as an Excel (.xlsx) workbook."
        ) from exc


def _find_observation_sheet(workbook: Workbook) -> Worksheet:
    for name in workbook.sheetnames:
        if name.strip().lower() == OBSERVATIONS_SHEET.lower():
            return workbook[name]
    # Fall back to the first sheet so simple single-sheet uploads work.
    first = workbook[workbook.sheetnames[0]]
    return first


def _read_header_map(sheet: Worksheet) -> dict[int, str]:
    header_map: dict[int, str] = {}
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    for index, cell in enumerate(first_row):
        if _is_blank(cell):
            continue
        name = str(cell).strip().lower().replace(" ", "_").replace("-", "_")
        name = _HEADER_ALIASES.get(name, name)
        if name in OBSERVATION_COLUMNS:
            header_map[index] = name
    return header_map


def _row_to_observation(values: dict[str, Any]) -> DailyObservation:
    observation_date = _parse_date(values["date"]) if "date" in values else None
    if observation_date is None:
        raise ValueError("Missing date.")

    temperature = values.get("temperature")
    if temperature is not None:
        try:
            temperature = float(temperature)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Temperature {temperature!r} is not a number.") from exc

    unit = values.get("temperature_unit")
    temperature_unit = _parse_enum(TemperatureUnit, unit, "temperature_unit") if unit else None

    temperature_time = _parse_time(values.get("temperature_time"))

    fluid = _build_fluid(values)
    cervix = _build_cervix(values)

    bleeding_raw = values.get("bleeding")
    bleeding = (
        _parse_enum(BleedingLevel, bleeding_raw, "bleeding") if bleeding_raw else BleedingLevel.NONE
    )

    return DailyObservation(
        observation_date=observation_date,
        waking_temperature=temperature,
        temperature_unit=temperature_unit,
        temperature_time=temperature_time,
        temperature_disturbed=_parse_bool(values.get("temperature_disturbed"), default=False),
        fluid=fluid,
        cervical_position=cervix,
        bleeding=bleeding,
        notes=str(values.get("notes") or ""),
    )


def _build_fluid(values: dict[str, Any]) -> FluidObservation | None:
    keys = (
        "fluid_sensation",
        "fluid_quantity",
        "fluid_color",
        "fluid_texture",
        "fluid_amount",
        "peak_quality",
    )
    if not any(key in values for key in keys):
        return None

    amount = values.get("fluid_amount")
    if amount is not None:
        try:
            amount = int(amount)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Fluid amount {amount!r} is not a whole number.") from exc

    sensation_raw = values.get("fluid_sensation")
    sensation = (
        _parse_enum(FluidSensation, sensation_raw, "fluid_sensation")
        if sensation_raw
        else FluidSensation.DRY
    )

    return FluidObservation(
        sensation=sensation,
        quantity=_parse_enum(FluidQuantity, values.get("fluid_quantity"), "fluid_quantity")
        if values.get("fluid_quantity")
        else FluidQuantity.NONE,
        color=_parse_enum(MucusColor, values.get("fluid_color"), "fluid_color")
        if values.get("fluid_color")
        else MucusColor.NONE,
        texture=_parse_enum(MucusTexture, values.get("fluid_texture"), "fluid_texture")
        if values.get("fluid_texture")
        else MucusTexture.NONE,
        amount=amount,
        peak_quality=_parse_bool(values.get("peak_quality"), default=False),
    )


def _build_cervix(values: dict[str, Any]) -> CervicalPositionObservation | None:
    keys = ("cervix_height", "cervix_firmness", "cervix_opening")
    present = [key for key in keys if key in values]
    if not present:
        return None
    if len(present) != len(keys):
        missing = ", ".join(sorted(set(keys) - set(present)))
        raise ValueError(f"Cervix observation needs height, firmness, and opening (missing: {missing}).")
    return CervicalPositionObservation(
        height=_parse_enum(CervixHeight, values["cervix_height"], "cervix_height"),
        firmness=_parse_enum(CervixFirmness, values["cervix_firmness"], "cervix_firmness"),
        opening=_parse_enum(CervixOpening, values["cervix_opening"], "cervix_opening"),
    )


def _write_observation_sheet(sheet: Worksheet, observations: Iterable[DailyObservation]) -> None:
    sheet.append(list(OBSERVATION_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for observation in sorted(observations, key=lambda item: item.observation_date):
        fluid = observation.fluid
        cervix = observation.cervical_position
        sheet.append(
            [
                observation.observation_date.isoformat(),
                observation.waking_temperature,
                observation.temperature_unit.value if observation.temperature_unit else None,
                observation.temperature_time.strftime("%H:%M") if observation.temperature_time else None,
                observation.temperature_disturbed or None,
                fluid.sensation.value if fluid else None,
                fluid.quantity.value if fluid and fluid.quantity != FluidQuantity.NONE else None,
                fluid.color.value if fluid and fluid.color != MucusColor.NONE else None,
                fluid.texture.value if fluid and fluid.texture != MucusTexture.NONE else None,
                fluid.amount if fluid else None,
                fluid.peak_quality or None if fluid else None,
                cervix.height.value if cervix else None,
                cervix.firmness.value if cervix else None,
                cervix.opening.value if cervix else None,
                observation.bleeding.value if observation.bleeding != BleedingLevel.NONE else None,
                observation.notes or None,
            ]
        )

    for index, name in enumerate(OBSERVATION_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(name) + 2, 12)
    sheet.freeze_panes = "A2"


def _write_settings_sheet(sheet: Worksheet, settings: AppSettings) -> None:
    sheet.append(["setting", "value"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for key, value in settings.as_dict().items():
        sheet.append([key, value])
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 18


def _parse_date(value: Any) -> date | None:
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_iso_date(str(value).strip())


def _parse_time(value: Any) -> time | None:
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    text = str(value).strip()
    if len(text) > 5 and text.count(":") == 2:
        text = text[: text.rfind(":")]
    return parse_hhmm_time(text)


def _parse_bool(value: Any, *, default: bool) -> bool:
    if _is_blank(value):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in _TRUE_STRINGS:
        return True
    if text in _FALSE_STRINGS:
        return False
    raise ValueError(f"Cannot interpret {value!r} as yes/no.")


def _parse_enum(enum_class: type, value: Any, column: str):
    text = str(value).strip().lower().replace(" ", "_")
    try:
        return enum_class(text)
    except ValueError as exc:
        allowed = ", ".join(item.value for item in enum_class)
        raise ValueError(f"{column} value {value!r} is not one of: {allowed}.") from exc


def _string_or(value: Any, default: str) -> str:
    if _is_blank(value):
        return default
    return str(value).strip().lower()


def _time_string_or(value: Any, default: str) -> str:
    if _is_blank(value):
        return default
    if isinstance(value, (datetime, time)):
        return value.strftime("%H:%M")
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
